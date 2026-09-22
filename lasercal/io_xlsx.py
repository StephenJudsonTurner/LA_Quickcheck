"""Reading and writing every XLSX file the MATLAB tool uses, byte-for-byte
compatible in layout so a data folder can be shared between the two tools.

Files (all in the data folder unless noted)
-------------------------------------------
<data>.xlsx              raw traces: Time + one column per isotope (7Li, 9Be, ...)
RunOrder.xlsx            one sample name per row, no header, column A
Intervals.xlsx           Row | back_start1 | back_stop1 | signal_start | signal_stop |
                         back_start2 | back_stop2 | CaO (wt%) | SiO2 (wt%) | Al2O3 (wt%)
DriftSelections.xlsx     Row | Standard | Interpolation | Exclusions
CalibrationSelections.xlsx  Row | NormElement | ForceInterceptZero | StandardSet | CalExclusions
DriftCorrected_All.xlsx  sheets LT_Corr_Data, LT_Corr_Data_CaNorm
ReducedDataExport.xlsx   sheets FinalResult, Signal_SEM_Norm, StandardValues, Intervals,
                         DriftSelections, CalibrationSelections
LAICPMS_Standard_Values.xlsx  (next to the program) standards table, first header blank

MATLAB writes strings as text cells, logicals as boolean cells, numbers as
numbers and NaN as empty cells; this module does the same.
"""
from __future__ import annotations

import math
import os
from typing import Dict, List, Optional, Sequence, Tuple

import numpy as np
import openpyxl
import pandas as pd

from .names import mangle_columns, make_valid_name


# ----------------------------------------------------------------------------
# Low-level sheet access
# ----------------------------------------------------------------------------
def sheet_names(path: str) -> List[str]:
    wb = openpyxl.load_workbook(path, read_only=True)
    try:
        return [ws.title for ws in wb.worksheets]
    finally:
        wb.close()


def read_sheet_values(path: str, sheet=None) -> List[List]:
    """All cell values of a sheet as a list of rows (None for blanks).

    ``sheet`` may be a name, a 0-based index or None (first sheet, like
    MATLAB's readtable default).
    """
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        if sheet is None:
            ws = wb.worksheets[0]
        elif isinstance(sheet, int):
            ws = wb.worksheets[sheet]
        else:
            ws = wb[sheet]
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
    finally:
        wb.close()
    # strip fully-empty trailing rows / columns
    while rows and all(v is None for v in rows[-1]):
        rows.pop()
    if rows:
        ncol = max(len(r) for r in rows)
        while ncol > 0 and all((len(r) < ncol or r[ncol - 1] is None) for r in rows):
            ncol -= 1
        rows = [list(r[:ncol]) + [None] * (ncol - len(r)) for r in rows]
    return rows


def _to_float(v) -> float:
    if v is None:
        return math.nan
    if isinstance(v, bool):
        return float(v)
    if isinstance(v, (int, float, np.integer, np.floating)):
        return float(v)
    try:
        return float(str(v).strip())
    except ValueError:
        return math.nan


def _numeric_frame(header: Sequence[str], rows: List[List]) -> pd.DataFrame:
    data = {h: [_to_float(r[i]) if i < len(r) else math.nan for r in rows] for i, h in enumerate(header)}
    return pd.DataFrame(data, columns=list(header))


# ----------------------------------------------------------------------------
# Readers
# ----------------------------------------------------------------------------
def read_run_order(path: str) -> List[str]:
    """RunOrder.xlsx read as MATLAB does: column A from row 1, as text."""
    rows = read_sheet_values(path)
    out = []
    for r in rows:
        v = r[0] if r else None
        if v is None:
            out.append('')
        elif isinstance(v, float) and v.is_integer():
            out.append(str(int(v)))
        else:
            out.append(str(v))
    # MATLAB's import keeps every row up to the last non-empty one
    while out and out[-1] == '':
        out.pop()
    return out


def read_raw_data(path: str, sheet=None) -> pd.DataFrame:
    """The batch trace file. Column names are mangled like MATLAB readtable
    ('7Li' -> 'x7Li'); all columns are numeric."""
    rows = read_sheet_values(path, sheet)
    if not rows:
        raise ValueError(f'{path} is empty.')
    header = mangle_columns(rows[0])
    df = _numeric_frame(header, rows[1:])
    return df


def read_standards(path: str, sheet=None) -> pd.DataFrame:
    """LAICPMS_Standard_Values.xlsx -> DataFrame with MATLAB column names
    ('Var1' for the blank first header, 'x7Li', ...)."""
    rows = read_sheet_values(path, sheet)
    header = mangle_columns(rows[0])
    body = rows[1:]
    data: Dict[str, list] = {}
    for i, h in enumerate(header):
        col = [r[i] if i < len(r) else None for r in body]
        if i == 0:
            data[h] = ['' if v is None else str(v) for v in col]
        else:
            data[h] = [_to_float(v) for v in col]
    return pd.DataFrame(data, columns=header)


def read_row_table(path: str, sheet=None, numeric: bool = False,
                   preserve_names: bool = True) -> pd.DataFrame:
    """A table written with WriteRowNames=true: first column = row names.

    numeric=True coerces every data column to float (Intervals.xlsx);
    otherwise cell values are kept as read (str / bool / number / None).
    """
    rows = read_sheet_values(path, sheet)
    if not rows:
        return pd.DataFrame()
    header = [str(h) if h is not None else '' for h in rows[0]]
    if not preserve_names:
        header = header[:1] + mangle_columns(header[1:])
    body = rows[1:]
    idx = [('' if r[0] is None else str(r[0])) for r in body]
    cols = header[1:]
    if numeric:
        df = _numeric_frame(cols, [r[1:] for r in body])
    else:
        df = pd.DataFrame([[r[i + 1] if i + 1 < len(r) else None for i in range(len(cols))] for r in body],
                          columns=cols, dtype=object)
    df.index = pd.Index(idx)
    df.index.name = None
    return df


def read_intervals(path: str, sheet=None) -> pd.DataFrame:
    return read_row_table(path, sheet, numeric=True)


# ----------------------------------------------------------------------------
# Writers
# ----------------------------------------------------------------------------
def _cell_value(v):
    """Convert a Python / numpy value to what MATLAB writetable would emit."""
    if v is None:
        return None
    if isinstance(v, (bool, np.bool_)):
        return bool(v)
    if isinstance(v, (np.integer,)):
        return int(v)
    if isinstance(v, (float, np.floating)):
        return None if math.isnan(float(v)) else float(v)
    if isinstance(v, (int,)):
        return v
    if isinstance(v, str):
        return v if v != '' else None
    try:
        if pd.isna(v):
            return None
    except (TypeError, ValueError):
        pass
    return str(v)


def _replace_sheet(wb: openpyxl.Workbook, name: str):
    """Return an empty sheet called `name`, keeping its position if it exists
    (MATLAB writetable replaces a sheet's content and keeps other sheets)."""
    if name in wb.sheetnames:
        pos = wb.sheetnames.index(name)
        wb.remove(wb[name])
        return wb.create_sheet(name, pos)
    return wb.create_sheet(name)


def _open_for_write(path: str) -> Tuple[openpyxl.Workbook, bool]:
    if os.path.isfile(path):
        return openpyxl.load_workbook(path), True
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    return wb, False


def _fill_sheet(ws, header: Sequence, rows: Sequence[Sequence]):
    ws.append([_cell_value(h) for h in header])
    for r in rows:
        ws.append([_cell_value(v) for v in r])


def write_row_table(path: str, df: pd.DataFrame, sheet: str = 'Sheet1',
                    row_label: str = 'Row', header_override: Optional[Sequence[str]] = None):
    """writetable(T, path, 'WriteRowNames', true, 'Sheet', sheet)."""
    wb, existed = _open_for_write(path)
    ws = _replace_sheet(wb, sheet)
    header = [row_label] + (list(header_override) if header_override is not None else list(df.columns))
    rows = [[idx] + list(df.iloc[i].tolist()) for i, idx in enumerate(df.index)]
    _fill_sheet(ws, header, rows)
    wb.save(path)


def write_table(path: str, df: pd.DataFrame, sheet: str = 'Sheet1'):
    """writetable(T, path, 'Sheet', sheet) for a table without row names."""
    wb, existed = _open_for_write(path)
    ws = _replace_sheet(wb, sheet)
    _fill_sheet(ws, list(df.columns), [list(df.iloc[i].tolist()) for i in range(len(df))])
    wb.save(path)


def write_sheets(path: str, sheets: List[Tuple[str, pd.DataFrame, Optional[str], Optional[Sequence[str]]]]):
    """Write several sheets in one save.  Each tuple is
    (sheet name, DataFrame, row label or None, header override or None)."""
    wb, existed = _open_for_write(path)
    for name, df, row_label, header_override in sheets:
        ws = _replace_sheet(wb, name)
        cols = list(header_override) if header_override is not None else list(df.columns)
        if row_label is None:
            _fill_sheet(ws, cols, [list(df.iloc[i].tolist()) for i in range(len(df))])
        else:
            _fill_sheet(ws, [row_label] + cols,
                        [[idx] + list(df.iloc[i].tolist()) for i, idx in enumerate(df.index)])
    wb.save(path)


# ----------------------------------------------------------------------------
# Convenience: the specific files
# ----------------------------------------------------------------------------
def write_run_order(path: str, names: Sequence[str]):
    """RunOrder.xlsx: one name per row in column A, no header."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Sheet1'
    for n in names:
        ws.append([str(n)])
    wb.save(path)


def read_sample_windows(path: str) -> pd.DataFrame:
    """SampleWindows.xlsx: Row | start_row | stop_row | laser_on_row (1-based rows of the data file)."""
    df = read_row_table(path, numeric=True)
    for c in ('start_row', 'stop_row', 'laser_on_row'):
        if c not in df.columns:
            raise ValueError(f'SampleWindows.xlsx must have a "{c}" column')
    return df


def write_intervals(path: str, intervals: pd.DataFrame):
    write_row_table(path, intervals, 'Sheet1')


def write_drift_selections(path: str, prefs: pd.DataFrame):
    write_row_table(path, prefs[['Standard', 'Interpolation', 'Exclusions']], 'Sheet1')


def write_calibration_selections(path: str, prefs: pd.DataFrame):
    write_row_table(path, prefs[['NormElement', 'ForceInterceptZero', 'StandardSet', 'CalExclusions']], 'Sheet1')


def write_drift_corrected(path: str, lt_corr: np.ndarray, sig_vars: Sequence[str],
                          display_names: Sequence[str], ca_iso: Optional[str]):
    """exportCorrected(): LT_Corr_Data and (if Ca present) LT_Corr_Data_CaNorm."""
    eps = np.finfo(float).eps
    tbl = pd.DataFrame(lt_corr, columns=list(sig_vars), index=list(display_names))
    sheets = [('LT_Corr_Data', tbl, 'Row', None)]
    if ca_iso is not None and ca_iso in sig_vars:
        den = lt_corr[:, list(sig_vars).index(ca_iso)].copy()
        den[(den == 0) | np.isnan(den)] = eps
        tbl_n = pd.DataFrame(lt_corr / den[:, None], columns=list(sig_vars), index=list(display_names))
        sheets.append(('LT_Corr_Data_CaNorm', tbl_n, 'Row', None))
    write_sheets(path, sheets)


def write_reduced_export(path: str, final: np.ndarray, sig_vars: Sequence[str],
                         display_names: Sequence[str], first_in_times: np.ndarray,
                         sem: Optional[np.ndarray], std_vals: Optional[pd.DataFrame],
                         intervals: Optional[pd.DataFrame], drift_prefs: Optional[pd.DataFrame],
                         cal_prefs: Optional[pd.DataFrame]):
    """exportCalibration(): the six-sheet reduced data workbook."""
    n = final.shape[0]
    fr = pd.DataFrame(final, columns=list(sig_vars), index=list(display_names))
    fr.insert(0, 'Analysis_Start_Time', np.asarray(first_in_times, float))
    fr.insert(0, 'RunOrder', np.arange(1, n + 1))
    header = ['RunOrder', 'Analysis Start Time'] + list(sig_vars)   # writecell C1 fix-up
    sheets = [('FinalResult', fr, 'Row', header)]
    if sem is not None:
        sheets.append(('Signal_SEM_Norm', pd.DataFrame(sem, columns=list(sig_vars), index=list(display_names)), 'Row', None))
    if std_vals is not None:
        sheets.append(('StandardValues', std_vals, None, None))
    if intervals is not None:
        sheets.append(('Intervals', intervals, 'Row', None))
    if drift_prefs is not None:
        sheets.append(('DriftSelections', drift_prefs[['Standard', 'Interpolation', 'Exclusions']], 'Row', None))
    if cal_prefs is not None:
        sheets.append(('CalibrationSelections',
                       cal_prefs[['NormElement', 'ForceInterceptZero', 'StandardSet', 'CalExclusions']], 'Row', None))
    write_sheets(path, sheets)
